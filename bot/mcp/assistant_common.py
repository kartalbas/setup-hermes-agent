"""Shared plumbing for the assistant MCP servers.

Two servers use this: ``m365_assistant.py`` (Microsoft Graph) and
``google_assistant.py`` (Gmail, Google Calendar, Drive). Both are stdio MCP
servers started by the agent's gateway; both hold one refresh token for one
account; both hand documents back as text.

Kept deliberately small and dependency-free: the standard library for the
protocol and HTTP, ``pypdf`` and ``python-docx`` only for document text.
"""

from __future__ import annotations

import base64
import json
import os
import stat
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass, field
from typing import Any, Callable, Dict, List, Optional, Tuple

PROTOCOL_VERSION = "2025-06-18"
USER_AGENT = "hermes-assistant/0.1"


def log(msg: str) -> None:
    """stderr only: stdout is the MCP channel."""
    sys.stderr.write(f"{time.strftime('%Y-%m-%d %H:%M:%S')} {msg}\n")
    sys.stderr.flush()


# ---------------------------------------------------------------------------
# HTTP
# ---------------------------------------------------------------------------

class HttpError(Exception):
    def __init__(self, status: int, method: str, url: str, body: str):
        self.status, self.method, self.url, self.body = status, method, url, body
        super().__init__(f"{method} {url} -> HTTP {status}: {body[:600]}")


def http(method: str, url: str, *, headers: Optional[Dict[str, str]] = None,
         json_body: Any = None, data: Optional[bytes] = None,
         params: Optional[Dict[str, Any]] = None, timeout: int = 60,
         raw: bool = False) -> Any:
    """One HTTP request. JSON in, JSON out unless ``raw``.

    Errors are raised as HttpError with the server's body, because Graph and
    Google both put the reason there and the agent needs to read it.
    """
    if params:
        clean = {k: v for k, v in params.items() if v is not None and v != ""}
        if clean:
            url += ("&" if "?" in url else "?") + urllib.parse.urlencode(clean, doseq=True)
    hdrs = {"User-Agent": USER_AGENT, "Accept": "application/json"}
    if headers:
        hdrs.update(headers)
    body = data
    if json_body is not None:
        body = json.dumps(json_body).encode("utf-8")
        hdrs.setdefault("Content-Type", "application/json")
    req = urllib.request.Request(url, data=body, method=method, headers=hdrs)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            payload = resp.read()
            if raw:
                return payload, dict(resp.headers)
            if not payload:
                return {}
            ctype = resp.headers.get("Content-Type", "")
            if "json" in ctype:
                return json.loads(payload.decode("utf-8"))
            return payload
    except urllib.error.HTTPError as e:
        text = e.read().decode("utf-8", "replace")
        raise HttpError(e.code, method, url, text) from None


def form_post(url: str, fields: Dict[str, str], timeout: int = 60) -> Dict[str, Any]:
    """application/x-www-form-urlencoded POST, JSON answer (OAuth endpoints).

    OAuth error answers come with 4xx AND a JSON body that carries the reason;
    they are returned, not raised, so the device-code poll can read
    ``authorization_pending``.
    """
    data = urllib.parse.urlencode(fields).encode("utf-8")
    req = urllib.request.Request(url, data=data, method="POST", headers={
        "User-Agent": USER_AGENT, "Accept": "application/json",
        "Content-Type": "application/x-www-form-urlencoded"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        text = e.read().decode("utf-8", "replace")
        try:
            return json.loads(text)
        except ValueError:
            raise HttpError(e.code, "POST", url, text) from None


# ---------------------------------------------------------------------------
# Token store — one file, one account, mode 0600.
# ---------------------------------------------------------------------------

class TokenStore:
    def __init__(self, path: str):
        self.path = path
        self.data: Dict[str, Any] = {}
        if os.path.exists(path):
            with open(path, "r", encoding="utf-8") as f:
                self.data = json.load(f)

    @property
    def refresh_token(self) -> str:
        return str(self.data.get("refresh_token") or "")

    @property
    def account(self) -> str:
        return str(self.data.get("account") or "")

    def access_token_valid(self, margin: int = 120) -> bool:
        return bool(self.data.get("access_token")) and \
            float(self.data.get("expires_at", 0)) - margin > time.time()

    def save(self, **fields: Any) -> None:
        self.data.update(fields)
        d = os.path.dirname(self.path)
        if d:
            os.makedirs(d, mode=0o700, exist_ok=True)
        tmp = f"{self.path}.tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(self.data, f, indent=1)
        os.chmod(tmp, stat.S_IRUSR | stat.S_IWUSR)
        os.replace(tmp, self.path)


# ---------------------------------------------------------------------------
# Document text
# ---------------------------------------------------------------------------

TEXT_TYPES = {".txt", ".md", ".csv", ".json", ".log", ".ics", ".eml", ".html", ".htm", ".xml", ".yaml", ".yml"}


def extract_text(name: str, data: bytes, max_chars: int = 20000) -> Dict[str, Any]:
    """Text of a document by its file extension.

    pdf via pypdf, docx via python-docx, plain text by decoding. Anything else
    is reported as unsupported rather than guessed. A scanned PDF without a
    text layer comes back empty — that is said explicitly.
    """
    ext = os.path.splitext(name.lower())[1]
    text = ""
    note = ""
    if ext == ".pdf":
        try:
            from pypdf import PdfReader  # type: ignore
        except ImportError:
            raise RuntimeError("pypdf is not installed in the assistant's environment")
        import io
        reader = PdfReader(io.BytesIO(data))
        pages = [p.extract_text() or "" for p in reader.pages]
        text = "\n\n".join(pages).strip()
        note = f"{len(reader.pages)} pages"
        if not text:
            note += "; no text layer (scanned image) — OCR is not available"
    elif ext == ".docx":
        try:
            import docx  # type: ignore
        except ImportError:
            raise RuntimeError("python-docx is not installed in the assistant's environment")
        import io
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts).strip()
    elif ext in TEXT_TYPES or not ext:
        text = data.decode("utf-8", "replace")
    else:
        return {"name": name, "supported": False, "size": len(data),
                "note": f"no text extraction for '{ext}' files (pdf, docx and plain text are supported)"}
    truncated = len(text) > max_chars
    return {"name": name, "supported": True, "size": len(data), "chars": len(text),
            "truncated": truncated, "note": note, "text": text[:max_chars]}


def b64(data: bytes) -> str:
    return base64.b64encode(data).decode("ascii")


def unb64(s: str) -> bytes:
    return base64.b64decode(s)


# ---------------------------------------------------------------------------
# MCP over stdio — the subset the gateway uses.
# ---------------------------------------------------------------------------

@dataclass
class Tool:
    name: str
    description: str
    schema: Dict[str, Any]
    fn: Callable[..., Any]

    def spec(self) -> Dict[str, Any]:
        schema = dict(self.schema)
        schema.setdefault("type", "object")
        schema.setdefault("additionalProperties", False)
        return {"name": self.name, "description": self.description, "inputSchema": schema}


@dataclass
class McpServer:
    name: str
    version: str
    instructions: str = ""
    tools: List[Tool] = field(default_factory=list)

    def tool(self, name: str, description: str, schema: Dict[str, Any]):
        def deco(fn):
            self.tools.append(Tool(name, description, schema, fn))
            return fn
        return deco

    # -- dispatch -----------------------------------------------------------
    def handle(self, msg: Dict[str, Any]) -> Optional[Dict[str, Any]]:
        method = msg.get("method")
        msg_id = msg.get("id")
        params = msg.get("params") or {}
        if method == "initialize":
            return self._ok(msg_id, {
                "protocolVersion": params.get("protocolVersion") or PROTOCOL_VERSION,
                "capabilities": {"tools": {"listChanged": False}},
                "serverInfo": {"name": self.name, "version": self.version},
                "instructions": self.instructions,
            })
        if method == "notifications/initialized" or msg_id is None:
            return None
        if method == "ping":
            return self._ok(msg_id, {})
        if method == "tools/list":
            return self._ok(msg_id, {"tools": [t.spec() for t in self.tools]})
        if method == "tools/call":
            return self._call(msg_id, params)
        return self._err(msg_id, -32601, f"method not found: {method}")

    def _call(self, msg_id: Any, params: Dict[str, Any]) -> Dict[str, Any]:
        name = params.get("name")
        args = params.get("arguments") or {}
        tool = next((t for t in self.tools if t.name == name), None)
        if tool is None:
            return self._err(msg_id, -32602, f"unknown tool: {name}")
        try:
            problems = validate(args, tool.spec()["inputSchema"])
            if problems:
                return self._ok(msg_id, self._content(f"invalid arguments: {'; '.join(problems)}", True))
            result = tool.fn(**args)
            text = result if isinstance(result, str) else json.dumps(result, ensure_ascii=False, indent=1, default=str)
            return self._ok(msg_id, self._content(text, False))
        except HttpError as e:
            return self._ok(msg_id, self._content(f"{e.method} {e.url} failed with HTTP {e.status}: {e.body[:1500]}", True))
        except Exception as e:  # noqa: BLE001 — reported to the model, never swallowed
            log(f"tool {name} failed: {e!r}")
            return self._ok(msg_id, self._content(f"{type(e).__name__}: {e}", True))

    @staticmethod
    def _content(text: str, is_error: bool) -> Dict[str, Any]:
        return {"content": [{"type": "text", "text": text}], "isError": is_error}

    @staticmethod
    def _ok(msg_id: Any, result: Any) -> Dict[str, Any]:
        return {"jsonrpc": "2.0", "id": msg_id, "result": result}

    @staticmethod
    def _err(msg_id: Any, code: int, message: str) -> Dict[str, Any]:
        return {"jsonrpc": "2.0", "id": msg_id, "error": {"code": code, "message": message}}

    # -- transport ----------------------------------------------------------
    def serve_stdio(self) -> None:
        log(f"{self.name} {self.version}: serving {len(self.tools)} tools on stdio")
        out = sys.stdout
        for line in sys.stdin:
            line = line.strip()
            if not line:
                continue
            try:
                msg = json.loads(line)
            except ValueError:
                out.write(json.dumps(self._err(None, -32700, "parse error")) + "\n")
                out.flush()
                continue
            reply = self.handle(msg)
            if reply is not None:
                out.write(json.dumps(reply, ensure_ascii=False) + "\n")
                out.flush()
        log(f"{self.name}: stdin closed, exiting")


# ---------------------------------------------------------------------------
# A small JSON-schema check — required keys, types, enums — enough to give the
# model a precise message instead of a stack trace.
# ---------------------------------------------------------------------------

_TYPES = {"string": str, "integer": int, "number": (int, float), "boolean": bool,
          "array": list, "object": dict}


def validate(args: Dict[str, Any], schema: Dict[str, Any], path: str = "") -> List[str]:
    problems: List[str] = []
    if not isinstance(args, dict):
        return [f"{path or 'arguments'} must be an object"]
    props = schema.get("properties") or {}
    for key in schema.get("required") or []:
        if key not in args:
            problems.append(f"missing required '{path}{key}'")
    if schema.get("additionalProperties") is False:
        for key in args:
            if key not in props:
                problems.append(f"unknown argument '{path}{key}'")
    for key, val in args.items():
        spec = props.get(key)
        if not spec or val is None:
            continue
        want = spec.get("type")
        if want and want in _TYPES:
            ok = isinstance(val, _TYPES[want])
            if want == "integer" and isinstance(val, bool):
                ok = False
            if want == "number" and isinstance(val, bool):
                ok = False
            if not ok:
                problems.append(f"'{path}{key}' must be {want}")
                continue
        if "enum" in spec and val not in spec["enum"]:
            problems.append(f"'{path}{key}' must be one of {spec['enum']}")
        if want == "array" and "items" in spec:
            item_spec = spec["items"]
            for i, item in enumerate(val):
                if item_spec.get("type") == "object":
                    problems += validate(item, item_spec, f"{path}{key}[{i}].")
                elif item_spec.get("type") in _TYPES and not isinstance(item, _TYPES[item_spec["type"]]):
                    problems.append(f"'{path}{key}[{i}]' must be {item_spec['type']}")
        if want == "object" and "properties" in spec:
            problems += validate(val, spec, f"{path}{key}.")
    return problems


def env_required(name: str) -> str:
    v = os.environ.get(name, "").strip()
    if not v:
        raise SystemExit(f"{name} is not set; the installer passes it in the MCP server's env")
    return v


# ---------------------------------------------------------------------------
# Worlds: private and business apart, below one root folder — the same rule
# in every drive the assistants write to.
# ---------------------------------------------------------------------------

def parse_worlds(value: str) -> Dict[str, str]:
    """'Business=_bus,Private=_pri' -> {'Business': '_bus', 'Private': '_pri'}: the
    folders directly under the root folder, each with the suffix its files carry."""
    out: Dict[str, str] = {}
    for part in value.replace(";", ",").split(","):
        if "=" in part:
            name, suffix = part.split("=", 1)
            if name.strip() and suffix.strip():
                out[name.strip()] = suffix.strip()
    return out


def _split_name(name: str) -> Tuple[str, str]:
    stem, dot, ext = name.rpartition(".")
    return (stem, "." + ext) if dot and stem else (name, "")


def world_of(path: str, root: str, worlds: Dict[str, str]) -> Optional[Tuple[str, List[str]]]:
    """(world name, path segments below the root) for a path under ROOT, else None."""
    if not worlds or not root:
        return None
    parts = [p for p in path.strip("/").split("/") if p]
    rootparts = [p for p in root.strip("/").split("/") if p]
    if len(parts) <= len(rootparts) or [p.lower() for p in parts[:len(rootparts)]] != [p.lower() for p in rootparts]:
        return None
    below = parts[len(rootparts):]
    match = next((w for w in worlds if w.lower() == below[0].lower()), None)
    if match is None:
        raise ValueError(f"'{path}' is under {root}/ but not in one of its worlds ("
                         + ", ".join(f"{root}/{w}/" for w in worlds) + f"); file it as e.g. {root}/{next(iter(worlds))}/" + "/".join(below))
    return match, below[1:]


def world_check(path: str, root: str, worlds: Dict[str, str], is_folder: bool = False) -> Optional[str]:
    """The private/business split, enforced where files are written: below the
    root folder the first segment names a world, and a file's name ends with
    that world's suffix before the extension. Returns the world, or None when
    the path is outside the root; raises ValueError with the corrected name."""
    found = world_of(path, root, worlds)
    if found is None:
        return None
    world, rest = found
    if is_folder or not rest:
        return world
    stem, ext = _split_name(rest[-1])
    suffix = worlds[world]
    if not stem.endswith(suffix):
        raise ValueError(f"file names under {root}/{world}/ end with '{suffix}' before the extension: use '{stem}{suffix}{ext}'")
    return world


def plan_world_targets(files: List[str], root: str, worlds: Dict[str, str], default: Optional[str] = None) -> List[Tuple[str, str]]:
    """Where existing files below the root (paths relative to it, not yet in a
    world) go: into DEFAULT (the first world unless given), or into the world
    that a folder name or a word in the file name already carries ("privat");
    the file name gets the suffix unless it has it. Returns (source, target)
    pairs with full paths."""
    if not worlds:
        return []
    names = list(worlds)
    default = next((w for w in names if default and w.lower() == default.lower()), names[0])
    plan = []
    for rel in files:
        parts = [p for p in rel.strip("/").split("/") if p]
        if not parts:
            continue
        world, hit_at = default, None
        for i, seg in enumerate(parts[:-1]):
            hit = next((w for w in names if w.lower() == seg.lower() or seg.lower().startswith(w.lower()[:4])), None)
            if hit:
                world, hit_at = hit, i
                break
        stem, ext = _split_name(parts[-1])
        if hit_at is None:
            words = stem.lower().replace("_", " ").replace("-", " ").split()
            hint = next((w for w in names if any(x.startswith(w.lower()[:4]) for x in words)), None)
            if hint:
                world = hint
        name = parts[-1] if stem.endswith(worlds[world]) else f"{stem}{worlds[world]}{ext}"
        folders = [p for i, p in enumerate(parts[:-1]) if i != hit_at]     # the segment that named the world is the world
        plan.append((f"{root}/{rel.strip('/')}", "/".join([root, world] + folders + [name])))
    return plan


# ---------------------------------------------------------------------------
# Companion text: every document filed below the root gets a Markdown twin
# with the same name holding the recognized text — searchable, quotable,
# readable without opening the scan.
# ---------------------------------------------------------------------------

DOCUMENT_EXTS = {".pdf", ".docx", ".doc", ".jpg", ".jpeg", ".png", ".heic", ".heif", ".webp", ".tif", ".tiff", ".gif", ".bmp"}
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".heic", ".heif", ".webp", ".tif", ".tiff", ".gif", ".bmp"}


def _ext(name: str) -> str:
    return ("." + name.rsplit(".", 1)[1].lower()) if "." in name.rsplit("/", 1)[-1] else ""


def is_document(name: str) -> bool:
    """A scan, photo, PDF or Word file — something whose text is not the file."""
    return _ext(name) in DOCUMENT_EXTS


def companion_path(path: str) -> str:
    """The twin's path: same folder, same stem, `.md`."""
    p = path.rstrip("/")
    stem = p.rsplit(".", 1)[0] if "." in p.rsplit("/", 1)[-1] else p
    return stem + ".md"


def recognized_text(name: str, data: Optional[bytes], text_md: Optional[str]) -> str:
    """The text that goes into the twin: what the caller recognized (text_md),
    or, for a PDF or Word file with a text layer, what the file itself yields.
    A photo or a scan without a text layer needs the caller's eyes — refused
    with the reason, so nothing is filed without its text."""
    if text_md and text_md.strip():
        return text_md.strip()
    ext = _ext(name)
    if ext in IMAGE_EXTS or data is None:
        raise ValueError(f"'{name}' is filed together with its recognized text: read the image first and pass it as text_md")
    got = extract_text(name, data, max_chars=200000)
    text = (got.get("text") or "").strip() if isinstance(got, dict) else ""
    if not text:
        raise ValueError(f"'{name}' has no text layer to extract: read it and pass the text as text_md")
    return text


def companion_markdown(name: str, world: Optional[str], text: str) -> str:
    """The twin's content: a small header, then the recognized text."""
    head = [f"# {name}", "", f"- source: {name}", f"- filed: {time.strftime('%Y-%m-%dT%H:%M:%S%z')}"]
    if world:
        head.append(f"- world: {world}")
    return "\n".join(head) + "\n\n" + text.rstrip() + "\n"
